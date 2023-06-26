from flask import Flask, request, render_template, send_file, session, flash
from flask_socketio import SocketIO, emit
from werkzeug.utils import secure_filename
import openai
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
from tenacity import retry, stop_after_attempt, wait_fixed
import time
import re
import docx2txt
import os
import uuid
import logging

app = Flask(__name__)
socketio = SocketIO(app)  # Create a new SocketIO instance
app.secret_key = os.urandom(24)  # It is required for session

logging.basicConfig(filename='app.log', filemode='w', format='%(name)s - %(levelname)s - %(message)s')

API_KEY = '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'  # Better to set it as environment variable
openai.api_key = API_KEY

UPLOAD_FOLDER = os.path.join(os.getcwd(), 'Uploaded_files')
OUTPUT_FOLDER = os.path.join(os.getcwd(), 'Output_files')
TEMPLATES_FOLDER = os.path.join(os.getcwd(), 'templates')


if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

if not os.path.exists(OUTPUT_FOLDER):
    os.makedirs(OUTPUT_FOLDER)


@retry(stop=stop_after_attempt(3), wait=wait_fixed(1))
def chatgpt_conversation(conversation_log, model_id):
    response = openai.ChatCompletion.create(
        model=model_id,
        temperature=0,
        messages=conversation_log
    )
    conversation_log.append({
        'role': response.choices[0].message.role,
        'content': response.choices[0].message.content.strip()
    })

    return conversation_log

@retry(stop=stop_after_attempt(3), wait=wait_fixed(1))
def process_answer(prompt, text_chunks, model_id):
    desired_outputs = []
    for text in text_chunks:
        # Clear conversations for each value
        conversations = []

        # Giving the prompt
        conversations.append({'role': 'user', 'content': prompt})
        conversations = chatgpt_conversation(conversations, model_id)

        # Giving the text
        conversations.append({'role': 'user', 'content': text})
        conversations = chatgpt_conversation(conversations, model_id)

        desired_output = conversations[-1]['content'].strip()
        desired_outputs.append(desired_output)

        time.sleep(1)  # Add a 1-second delay between requests

    abbreviations = []
    full_names = []
    for output in desired_outputs:
        output = re.sub(r'^---.*$', '', output, flags=re.MULTILINE)  
        output = re.sub(r'^Please note.*$', '', output, flags=re.MULTILINE)  
        output = re.sub(r'^Based on.*$', '', output, flags=re.MULTILINE)
        output = re.sub(r'^Note:.*$', '', output, flags=re.MULTILINE)  
        
        output_lines = output.split('\n')
        for line in output_lines:
            if '|' in line: 
                abbr, full_name = line.split(' | ')
                abbreviations.append(abbr.strip())
                full_names.append(full_name.strip())

    abbr_fullname_dict = {abbr: fullname for abbr, fullname in zip(abbreviations, full_names)}

    abbreviations_to_exclude = ['DET','Abbreviation']
    full_names_to_exclude = ['Data Extraction','Possible Full Versions']
    for abbreviation, full_name in zip(abbreviations_to_exclude, full_names_to_exclude):
        if abbreviation in abbr_fullname_dict:
            del abbr_fullname_dict[abbreviation]
        if full_name in abbr_fullname_dict.values():
            del abbr_fullname_dict[abbreviation]

    sorted_abbr_fullname = sorted(abbr_fullname_dict.items())
    abbreviations, full_names = zip(*sorted_abbr_fullname)

    return abbreviations, full_names

def chunks(lst, n):
    for i in range(0, len(lst), n):
        yield lst[i:i + n]

    
@app.route('/', methods=['GET', 'POST'])
def upload_file():
    if request.method == 'POST':
        try:
            file = request.files['file']   
            if file and allowed_file(file.filename):  # check if the post request has the file part
                filename = secure_filename(file.filename)
                session_id = str(uuid.uuid4())  # generating a unique session id
                filename_base, file_extension = os.path.splitext(filename)  # split filename and extension
                filename = f"{filename_base}_{session_id}{file_extension}"  # append the session id to the filename

                file.save(os.path.join(UPLOAD_FOLDER, filename))  # save the uploaded file

                # process the document and create the output file
                output_filename = process_file(filename)
                session['output_filename'] = output_filename  # storing output filename in session
                return render_template('upload.html', results="Document processed and table created successfully.")

        except Exception as e:
            logging.error(f"Error occurred: {e}")
            socketio.emit('error', {'error_message': str(e)}, namespace='/test')
            return render_template('error.html', error_message=str(e)), 500

        else:
            return render_template('error.html', error_message="Allowed file types are .doc, .docx"), 400

    return render_template('upload.html')

@app.route('/download')
def download_file():
    try:
        output_filename = session.get('output_filename', None)
        if output_filename:
            return send_file(os.path.join(OUTPUT_FOLDER, output_filename), as_attachment=True)
        else:
            return render_template('error.html', error_message="No file to download"), 400

    except Exception as e:
        logging.error(f"Error occurred: {e}")
        socketio.emit('error', {'error_message': str(e)}, namespace='/test')
        return render_template('error.html', error_message=str(e)), 500
        

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in {'doc', 'docx'}


def process_file(filename):
    try:
        # path of the file in the UPLOAD_FOLDER
        filepath_in_upload_folder = os.path.join(UPLOAD_FOLDER, filename)

        # Read and process the document text
        text = docx2txt.process(filepath_in_upload_folder)
 
        text = re.sub(r'\s+', ' ', text)  # Replace multiple whitespaces with a single one
        text_chunks = list(chunks(text.split(' '), 600))  # Split text into chunks of 600 words
        text_chunks = [' '.join(chunk) for chunk in text_chunks]  # Join words in each chunk
        
        # Derive output filename from input filename
        filename_base, file_extension = os.path.splitext(filename)  # This gets the filename without the extension
        output_filename = f"{filename_base}_abbreviation.docx"
        output_path = os.path.join(OUTPUT_FOLDER, output_filename)  # full path for output file
    
        # Load the template document
        template = os.path.join(TEMPLATES_FOLDER, "Abbretemplate.docx") # Replace with your actual template filename
        doc = Document(template)

        # Define the prompt and process text chunks
        prompt = "I want your assistance in making an abbreviation or acronym list from a medical report. Abbreviations and acronyms are defined as words that contain 2 or more letters that may or may not be fully capitalized. The abbreviated forms may or may not be in the text next to their corresponding full form in brackets. Finally, there may be abbreviations that are not defined within the text. Please list those as well by looking up what they stand for. If you locate more than one possible interpretation for an abbreviation, please list them separately. Additionally, the abbreviations should be presented in alphabetical order.  Give the output in this format always. Abbreviation | Full form. Example HRQoL | Health-Related Quality of Life; RCT | Randomised Controlled Trial. Abbreviation and Full form must be separated by this symbol'|' ###"
        # Process text chunks
        abbreviations, full_names = process_answer(prompt, text_chunks, 'gpt-3.5-turbo-0613')

        # Add the table to the end of the document
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'

        # Set table headers
        for i, name in enumerate(["Abbreviation", "Full Forms"]):
            cell = table.cell(0, i)
            cell.text = name
            cell._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="#D9D9D9"/>'.format(nsdecls('w'))))  # Set cell background color

        # Add data to table
        for abbr, full_name in zip(abbreviations, full_names):
            cells = table.add_row().cells
            cells[0].text = abbr
            cells[1].text = full_name

        # Save the document with the new filename
        doc.save(output_path)
        return output_filename  # return the output filename

    except Exception as e:
        logging.error(f"Error occurred: {e}")
        emit('error', {'error_message': str(e)})  # send error status update using socketio
        return None
  

if __name__ == '__main__':
    socketio.run(app, debug=True)  # use socketio.run instead of app.run
